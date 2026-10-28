"""Fix 28/10/2026, round-2 tester findings on the HTTP boundary and account rules.

R2-03 integers past what the database stores, R2-04 non-finite JSON numbers,
R2-05 body limits and echo-free 422s, R2-09 sessions after a deactivation,
R2-10 API docs off by default, R1-USERNAME case-insensitive usernames,
R1-AUDIT attempted usernames, R1-PHONE-EMAIL customer contacts, R1-AI-TEXT,
R1-EXCEL/R2-11 the lot report as Excel, A-09 catalog names, A-13/B-19 the
demo server's SPA shell. Every test runs against the in-memory test app.
"""
import csv
import importlib.util
import logging
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from core.config import Settings, settings
from models.audit_log import AuditLog
from models.customer import Customer
from models.role import Role
from models.user import User
from models.vehicle_type import VehicleType
from models.zone import Zone
from services.auth_service import AuthService
from test_expansion_sites import env  # noqa: F401

ROOT = Path(__file__).resolve().parents[1]
BIG = str(2**63)


def _role(db, name):
    row = db.scalar(select(Role).where(Role.name == name))
    if row is None:
        row = Role(name=name)
        db.add(row)
        db.flush()
    return row


def _user(db, username, role="customer", password="Password123"):
    user = User(username=username, full_name=f"Tài khoản {username}", role_id=_role(db, role).id,
                password_hash=AuthService.get_password_hash(password), is_active=True)
    db.add(user)
    db.commit()
    return user


def _headers(user):
    token = AuthService().create_access_token(user.id, user.username, user.role.name, password_hash=user.password_hash)
    return {"Authorization": f"Bearer {token}"}


def _login(client, username, password="Password123"):
    return client.post("/api/auth/login", json={"username": username, "password": password})


# --- R2-03 ---------------------------------------------------------------------------------

def test_integers_past_the_database_range_are_422_never_500(client, db_session):
    manager = _user(db_session, "range_manager", "manager")
    headers = _headers(manager)
    for path in (f"/api/v1/customers/{BIG}", f"/api/v1/customers?skip={BIG}", f"/api/v1/vehicles/{BIG}",
                 f"/dashboard/recent-sessions?limit={BIG}", f"/api/v2/sites/{BIG}/sessions"):
        response = client.get(path, headers=headers)
        assert response.status_code == 422, (path, response.status_code, response.text)
        assert response.json()["detail"], path


def test_paging_past_the_database_range_is_422_on_site_and_legacy_lists(client, db_session):
    from expansion.site_models import ParkingSite, SiteMembership
    manager = _user(db_session, "paging_manager", "manager")
    site = ParkingSite(name="Bãi phân trang")
    db_session.add(site)
    db_session.flush()
    db_session.add(SiteMembership(site_id=site.id, user_id=manager.id, role="manager"))
    db_session.commit()
    headers = _headers(manager)
    for path in (f"/api/v2/sites/{site.id}/sessions?offset={BIG}", f"/api/v2/sites/{site.id}/payments?page={2**63 - 1}",
                 f"/api/v2/sites/{site.id}/payments?page={BIG}", f"/api/v2/sites/{site.id}/refund-requests?offset={BIG}",
                 f"/parking/search?page={2**63 - 1}", f"/api/v1/parking-sessions?skip={BIG}", f"/api/v1/monthly-passes?skip={BIG}"):
        response = client.get(path, headers=headers)
        assert response.status_code == 422, (path, response.status_code, response.text)


def test_an_out_of_range_number_is_logged_without_query_values(client, db_session, caplog):
    """Round 2 (G5): the 422 stays, but a WARNING names the method, path and exception, so an
    arithmetic overflow inside the server stays visible; query values never reach the log."""
    headers = _headers(_user(db_session, "overflow_manager", "manager"))
    with caplog.at_level(logging.WARNING, logger="main"):
        response = client.get(f"/api/v1/customers?skip={BIG}&search=0901234567", headers=headers)
    assert response.status_code == 422
    records = [record for record in caplog.records if record.name == "main" and record.levelno == logging.WARNING]
    assert len(records) == 1, [record.getMessage() for record in caplog.records]
    message = records[0].getMessage()
    assert "GET /api/v1/customers" in message and "OverflowError" in message
    assert BIG not in message and "0901234567" not in message


def test_recent_sessions_limit_has_an_upper_bound(client, db_session):
    headers = _headers(_user(db_session, "recent_manager", "manager"))
    assert client.get("/dashboard/recent-sessions?limit=100", headers=headers).status_code == 200
    assert client.get("/dashboard/recent-sessions?limit=101", headers=headers).status_code == 422
    assert client.get("/dashboard/recent-sessions?limit=100000", headers=headers).status_code == 422


# --- R2-04 ---------------------------------------------------------------------------------

@pytest.mark.parametrize("raw", ['{"username": NaN, "password": "x"}', '{"username": "a", "password": Infinity}',
                                 '{"username": -Infinity, "password": "x"}', '{"username": 1e400, "password": "x"}'])
def test_non_finite_json_numbers_are_422_with_a_json_body(client, raw):
    response = client.post("/api/auth/login", content=raw, headers={"Content-Type": "application/json"})
    assert response.status_code == 422, response.text
    assert response.headers["content-type"].startswith("application/json")
    assert "NaN" in response.json()["detail"][0]["msg"]


def test_a_non_finite_price_never_reaches_the_catalog(client, db_session, vehicle_type):
    headers = {**_headers(_user(db_session, "price_manager", "manager")), "Content-Type": "application/json"}
    for value in ("Infinity", "-Infinity", "NaN", "1e400"):
        raw = '{"vehicle_type_id": %d, "ticket_type": "HOURLY", "price": %s, "effective_date": "2026-10-28"}' % (vehicle_type.id, value)
        response = client.post("/api/v1/price-configs", content=raw, headers=headers)
        assert response.status_code == 422, (value, response.text)


# --- R2-05 ---------------------------------------------------------------------------------

def test_an_oversized_body_gets_413_and_is_never_echoed(client):
    username = "u" * 10_000_000
    response = client.post("/api/auth/register", json={"username": username, "password": "Abcdefgh123!", "full_name": "x"})
    assert response.status_code == 413
    assert len(response.content) < 1000
    chunked = client.post("/api/auth/register", content=iter([b'{"username": "', b"u" * 2_000_000, b'"}']),
                          headers={"Content-Type": "application/json"})
    assert chunked.status_code == 413


def test_a_validation_error_never_echoes_the_submitted_value(client):
    username = "u" * 600_000
    response = client.post("/api/auth/register", json={"username": username, "password": "Abcdefgh123!", "full_name": "x"})
    assert response.status_code == 422
    assert len(response.content) < 2000 and "uuuuuuuuuu" not in response.text
    messages = [error["msg"] for error in response.json()["detail"]]
    assert "Tên đăng nhập: tối đa 50 ký tự." in messages and "Họ tên: cần ít nhất 2 ký tự." in messages
    assert all("input" not in error and "ctx" not in error for error in response.json()["detail"])


# --- R2-09 ---------------------------------------------------------------------------------

def test_reactivation_never_revives_a_token_from_before_the_deactivation(client, db_session):
    admin = _user(db_session, "tokens_admin", "admin")
    _user(db_session, "tokens_target")
    old = {"Authorization": "Bearer " + _login(client, "tokens_target").json()["access_token"]}
    target = db_session.scalar(select(User).where(User.username == "tokens_target"))
    assert client.put(f"/api/v1/users/{target.id}", json={"is_active": False}, headers=_headers(admin)).status_code == 200
    assert client.get("/api/auth/me", headers=old).status_code == 401
    assert client.put(f"/api/v1/users/{target.id}", json={"is_active": True}, headers=_headers(admin)).status_code == 200
    assert client.get("/api/auth/me", headers=old).status_code == 401
    fresh = _login(client, "tokens_target")
    assert fresh.status_code == 200, fresh.text
    new = {"Authorization": "Bearer " + fresh.json()["access_token"]}
    edited = client.put("/api/auth/me", json={"username": "tokens_target2", "full_name": "Tên mới"}, headers=new)
    assert edited.status_code == 200, edited.text
    assert client.get("/api/auth/me", headers=new).status_code == 200


# --- R2-10 ---------------------------------------------------------------------------------

def test_api_docs_are_off_unless_a_development_setting_turns_them_on(monkeypatch):
    from main import api_docs_urls
    monkeypatch.delenv("API_DOCS_ENABLED", raising=False)  # conftest turns them on for contract tests
    assert Settings(_env_file=None, SECRET_KEY="docs-test-secret").API_DOCS_ENABLED is False
    closed = TestClient(FastAPI(**api_docs_urls(False)))
    for path in ("/openapi.json", "/docs", "/redoc", "/docs/oauth2-redirect"):
        assert closed.get(path).status_code == 404, path
    assert TestClient(FastAPI(**api_docs_urls(True))).get("/openapi.json").status_code == 200


# --- R1-USERNAME and R1-AUDIT --------------------------------------------------------------

def test_usernames_are_unique_ignoring_case_on_every_path(client, db_session):
    admin = _user(db_session, "case_admin", "admin")
    _user(db_session, "customer_demo")
    other = _user(db_session, "case_customer")
    taken = "Tên đăng nhập đã tồn tại"
    response = client.post("/api/auth/register", json={"username": "CUSTOMER_DEMO", "password": "Password123", "full_name": "Khách"})
    assert response.status_code == 409 and response.json()["detail"] == taken
    response = client.put("/api/auth/me", json={"username": "Customer_Demo", "full_name": "Khách"}, headers=_headers(other))
    assert response.status_code == 409 and response.json()["detail"] == taken
    staff_role = _role(db_session, "staff")
    db_session.commit()
    response = client.post("/api/v1/users", headers=_headers(admin), json={"username": "CASE_ADMIN", "full_name": "Nhân viên",
        "password": "StaffPass1", "role_id": staff_role.id, "is_active": True})
    assert response.status_code == 409 and response.json()["detail"] == taken
    response = client.put(f"/api/v1/users/{other.id}", headers=_headers(admin), json={"username": "CUSTOMER_demo"})
    assert response.status_code == 409 and response.json()["detail"] == taken
    # A user may change the case of their own name, and a new name still registers.
    assert client.put("/api/auth/me", json={"username": "Case_Customer", "full_name": "Khách"}, headers=_headers(other)).status_code == 200
    assert client.post("/api/auth/register", json={"username": "brand_new", "password": "Password123", "full_name": "Khách mới"}).status_code == 201


def test_accounts_that_were_case_twins_before_the_rule_keep_saving_their_profile(client, db_session):
    """Round 2 (BS-2): only a new or changed name is checked, as on the admin path."""
    lower, upper = _user(db_session, "pair_user"), _user(db_session, "PAIR_USER")
    for user in (lower, upper):
        saved = client.put("/api/auth/me", json={"username": user.username, "full_name": "Tên mới"}, headers=_headers(user))
        assert saved.status_code == 200 and saved.json()["full_name"] == "Tên mới", saved.text
    # A changed name that collides with the twin is still refused.
    renamed = client.put("/api/auth/me", json={"username": "Pair_User", "full_name": "Tên mới"}, headers=_headers(lower))
    assert renamed.status_code == 409 and renamed.json()["detail"] == "Tên đăng nhập đã tồn tại"


def test_login_and_register_audit_rows_name_the_attempted_account(client, db_session):
    _user(db_session, "audited_user")
    assert _login(client, "audited_user", "wrong-password").status_code == 401
    assert _login(client, "audited_user").status_code == 200
    assert client.post("/api/auth/register", json={"username": "audited_new", "password": "Password123",
        "full_name": "Khách mới"}).status_code == 201
    rows = db_session.scalars(select(AuditLog).order_by(AuditLog.id)).all()
    assert [(row.action, row.username, row.success) for row in rows] == [
        ("LOGIN", "audited_user", False), ("LOGIN", "audited_user", True), ("REGISTER", "audited_new", True)]
    assert all(row.user_id is None for row in rows)
    assert not any("Password123" in str(value) or "wrong-password" in str(value)
                   for row in rows for value in vars(row).values())


# --- R1-PHONE-EMAIL ------------------------------------------------------------------------

def test_customer_phone_and_email_are_checked_when_new_and_seeded_rows_stay_editable(client, db_session):
    headers = _headers(_user(db_session, "contacts_manager", "manager"))
    created = client.post("/api/v1/customers", json={"full_name": "Khách", "phone_number": "abc"}, headers=headers)
    assert created.status_code == 422 and "Số điện thoại không hợp lệ" in created.json()["detail"][0]["msg"]
    assert client.post("/api/v1/customers", json={"full_name": "Khách", "phone_number": "0901 234.567"}, headers=headers).status_code == 201
    seeded = Customer(full_name="Khách DEMO - valid", phone_number="DEMO-CUSTOMER-1", email="demo-valid@example.invalid")
    db_session.add(seeded)
    db_session.commit()
    unchanged = client.put(f"/api/v1/customers/{seeded.id}", headers=headers,
        json={"full_name": "x", "phone_number": "DEMO-CUSTOMER-1", "email": "demo-valid@example.invalid"})
    assert unchanged.status_code == 200, unchanged.text
    bad_phone = client.put(f"/api/v1/customers/{seeded.id}", headers=headers, json={"phone_number": "abc"})
    assert bad_phone.status_code == 422 and bad_phone.json()["detail"].startswith("Số điện thoại không hợp lệ")
    bad_email = client.put(f"/api/v1/customers/{seeded.id}", headers=headers, json={"email": "new@example.invalid"})
    assert bad_email.status_code == 422 and bad_email.json()["detail"] == "Email không hợp lệ."
    fixed = client.put(f"/api/v1/customers/{seeded.id}", headers=headers, json={"email": "Demo@Example.com"})
    assert fixed.status_code == 200 and fixed.json()["email"] == "demo@example.com"


# --- R1-AI-TEXT ----------------------------------------------------------------------------

def test_ai_unavailable_text_names_no_configuration(client, db_session):
    headers = _headers(_user(db_session, "ai_manager", "manager"))
    for flag, key in ((False, "test-key"), (True, "")):
        with patch.object(settings, "AI_ENABLED", flag), patch.object(settings, "GEMINI_API_KEY", key):
            response = client.post("/ai/question", json={"question": "Hôm nay có bao nhiêu xe?"}, headers=headers)
        assert response.status_code == 503
        detail = response.json()["detail"]
        assert detail.startswith("Trợ lý AI hiện chưa được bật") and "AI_ENABLED" not in detail and "GEMINI" not in detail


# --- R1-EXCEL / R2-11 ----------------------------------------------------------------------

def test_lot_report_exports_the_same_rows_as_csv_and_excel(env):
    params = {"period": "week", "anchor_date": "2026-10-26"}
    url = f"/api/v2/sites/{env.a.id}/reports/export"
    as_csv = env.client.get(url, params=params)
    as_xlsx = env.client.get(url, params={**params, "format": "xlsx"})
    assert as_csv.status_code == 200 and as_xlsx.status_code == 200, as_xlsx.text
    assert as_xlsx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="parking-%d-week-2026-10-26.xlsx"' % env.a.id in as_xlsx.headers["content-disposition"]
    sheet = load_workbook(BytesIO(as_xlsx.content)).active
    excel = [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    text = [row for row in csv.reader(StringIO(as_csv.content.decode("utf-8-sig")))]
    # Two requests, two "as of" instants for the live availability rows; everything else is equal.
    def stable(rows):
        return [row for row in rows if row[1] != "Thời điểm"]
    assert len(excel) == len(text) and stable(excel) == stable(text)
    assert ["Thông tin", "Từ ngày", "", "", "", "2026-10-20", ""] in excel
    assert ["Thông tin", "Đến ngày", "", "", "", "2026-10-26", ""] in excel
    assert env.client.get(url, params={**params, "format": "pdf"}).status_code == 422


def test_excel_export_keeps_formulas_out(env):
    env.a.name = "=HYPERLINK(\"http://example.invalid\")"
    env.db.commit()
    response = env.client.get(f"/api/v2/sites/{env.a.id}/reports/export", params={"format": "xlsx"})
    sheet = load_workbook(BytesIO(response.content)).active
    names = [row[5] for row in sheet.iter_rows(values_only=True) if row[1] == "Bãi"]
    assert names == ["'=HYPERLINK(\"http://example.invalid\")"]
    assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)


# --- A-09 ----------------------------------------------------------------------------------

def test_catalog_names_collapse_spaces_and_ignore_accents(client, db_session):
    headers = _headers(_user(db_session, "names_manager", "manager"))
    db_session.add_all([VehicleType(name="Xe máy DEMO"), VehicleType(name="Ô tô DEMO")])
    db_session.commit()
    for name in ("Xe may DEMO", "  ô  TÔ   demo ", "xe MÁY demo"):
        response = client.post("/api/v1/vehicle-types", json={"name": name, "requires_plate": True}, headers=headers)
        assert response.status_code == 409 and response.json()["detail"] == "Tên loại xe đã tồn tại.", name
    spaced = client.post("/api/v1/vehicle-types", json={"name": "xe   xc", "requires_plate": True}, headers=headers)
    assert spaced.status_code == 201 and spaced.json()["name"] == "xe xc", spaced.text
    assert client.post("/api/v1/vehicle-types", json={"name": "xe xc", "requires_plate": True}, headers=headers).status_code == 409
    assert client.post("/api/v1/vehicle-types", json={"name": "Xe đạp điện", "requires_plate": False}, headers=headers).status_code == 201
    for name in ("Khu   Bắc", "khu bac"):
        response = client.post("/api/v1/zones", json={"name": name, "capacity": 5}, headers=headers)
        assert response.status_code == (201 if name == "Khu   Bắc" else 409), (name, response.text)
    assert db_session.scalar(select(Zone).where(Zone.name == "Khu Bắc")) is not None


def test_site_zone_names_collapse_spaces_and_ignore_accents(env):
    env.actor["user"] = env.staff
    zone = env.client.post(f"/api/v2/sites/{env.a.id}/zones", json={"name": "Khu   Đông", "capacity": 5})
    assert zone.status_code == 201 and zone.json()["name"] == "Khu Đông", zone.text
    again = env.client.post(f"/api/v2/sites/{env.a.id}/zones", json={"name": "khu dong", "capacity": 5})
    assert again.status_code == 409


def test_near_duplicate_rows_from_before_the_rule_stay_editable(client, db_session):
    """Round 2 (BD-1): only a changed name is checked; a row saved with its own name never conflicts."""
    headers = _headers(_user(db_session, "pairs_manager", "manager"))
    first_type, second_type = VehicleType(name="Xe máy"), VehicleType(name="Xe may")
    first_zone, second_zone, other_zone = Zone(name="Khu Đông", capacity=5), Zone(name="Khu Dong", capacity=5), Zone(name="Khu Tây", capacity=5)
    db_session.add_all([first_type, second_type, first_zone, second_zone, other_zone])
    db_session.commit()
    zone = client.put(f"/api/v1/zones/{second_zone.id}", headers=headers, json={"name": "Khu Dong", "capacity": 6, "is_active": True})
    assert zone.status_code == 200, zone.text
    kind = client.put(f"/api/v1/vehicle-types/{second_type.id}", headers=headers, json={"name": "Xe may", "is_active": False})
    assert kind.status_code == 200, kind.text
    assert client.put(f"/api/v1/vehicle-types/{first_type.id}", headers=headers, json={"name": "Xe Máy"}).status_code == 200
    # A changed name that matches another row is still refused.
    taken = client.put(f"/api/v1/zones/{second_zone.id}", headers=headers, json={"name": "khu tay"})
    assert taken.status_code == 409 and taken.json()["detail"] == "Tên khu vực đã tồn tại."


def test_a_near_duplicate_site_zone_stays_editable(env):
    env.db.add_all([Zone(name="Khu Đông", capacity=5, site_id=env.a.id), Zone(name="Khu Dong", capacity=5, site_id=env.a.id)])
    env.db.commit()
    second = env.db.scalar(select(Zone).where(Zone.name == "Khu Dong"))
    env.actor["user"] = env.staff
    patched = env.client.patch(f"/api/v2/sites/{env.a.id}/zones/{second.id}", json={"name": "khu  dong", "capacity": 7})
    assert patched.status_code == 200 and patched.json()["name"] == "khu dong", patched.text


def test_two_creates_of_the_same_name_are_serialized(tmp_path, monkeypatch):
    """Round 2 (BD-2): the check and the insert of a catalog name run under one catalog lock."""
    import threading
    from fastapi import HTTPException
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from crud import vehicle_type as crud_vt
    from database import Base
    from routers.vehicle_type import create_vehicle_type
    from schemas.vehicle_type import VehicleTypeCreate
    engine = create_engine(f"sqlite:///{(tmp_path / 'catalog.db').as_posix()}", connect_args={"timeout": 10})
    Base.metadata.create_all(engine)
    checked, real = threading.Event(), crud_vt.get_vehicle_type_by_name

    def slow_first_check(db, name, **kwargs):
        found = real(db, name, **kwargs)
        if threading.current_thread().name == "first":
            checked.set()
            threading.Event().wait(0.4)  # the second request arrives between check and insert
        return found
    monkeypatch.setattr(crud_vt, "get_vehicle_type_by_name", slow_first_check)
    results = {}

    def create(name):
        if threading.current_thread().name == "second":
            checked.wait(5)
        with Session(engine) as db:
            try:
                results[threading.current_thread().name] = create_vehicle_type(VehicleTypeCreate(name=name), db).name
            except HTTPException as exc:
                results[threading.current_thread().name] = exc.status_code
    threads = [threading.Thread(target=create, args=(name,), name=label) for label, name in (("first", "Xe máy"), ("second", "Xe may"))]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(15)
    engine.dispose()
    assert results == {"first": "Xe máy", "second": 409}


# --- A-13 / B-19 ---------------------------------------------------------------------------

def _demo_client(tmp_path, monkeypatch):
    spec = importlib.util.spec_from_file_location("parkingai_demo_server_hardening", ROOT / "scripts/demo_server.py")
    demo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(demo)
    dist = tmp_path / "frontend" / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<html>Demo SPA</html>")
    (dist / "assets" / "app.js").write_text("console.log(1)")
    monkeypatch.setattr(demo, "ROOT", tmp_path)
    monkeypatch.setattr("db_rollout.check_database_readiness", lambda *_args, **_kwargs: None)
    return TestClient(demo.build_demo_app())


def test_demo_spa_shell_carries_the_frontend_security_headers(tmp_path, monkeypatch):
    client = _demo_client(tmp_path, monkeypatch)
    for path in ("/", "/login", "/sites", "/config.js", "/assets/app.js"):
        headers = client.get(path).headers
        assert headers["x-content-type-options"] == "nosniff" and headers["x-frame-options"] == "DENY", path
        assert headers["strict-transport-security"].startswith("max-age=") and headers["referrer-policy"], path
        assert headers["permissions-policy"] == "camera=(self), microphone=(), geolocation=()", path
    # API answers keep their own, stricter policy.
    assert client.get("/api/auth/me").headers["permissions-policy"] == "camera=(), microphone=(), geolocation=()"


def test_demo_unknown_page_renders_the_spa_but_api_404s_stay_json(tmp_path, monkeypatch):
    client = _demo_client(tmp_path, monkeypatch)
    for path in ("/khong-ton-tai", "/portal/abc"):
        page = client.get(path, headers={"Accept": "text/html,application/xhtml+xml"})
        assert page.status_code == 200 and page.text == "<html>Demo SPA</html>", path
        assert page.headers["content-type"].startswith("text/html")
    api = client.get("/khong-ton-tai", headers={"Accept": "application/json"})
    assert api.status_code == 404 and api.json() == {"detail": "Not Found"}
    typed = client.get("/api/v2/khong-ton-tai", headers={"Accept": "text/html"})
    assert typed.status_code == 404 and typed.headers["content-type"].startswith("application/json")
