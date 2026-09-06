"""Ledger/shift regressions use only the isolated pytest database."""

from datetime import datetime, timedelta

import pytest
from fastapi import HTTPException
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError

from models.payment import Payment
from models.cash_shift import CashShift
from models.monthly_pass import MonthlyPass
from models.parking_session import ParkingSession
from models.role import Role
from models.user import User
from models.customer import Customer
from models.vehicle import Vehicle
from models.vehicle_type import VehicleType
from services.auth_service import AuthService
from services.payment_service import PaymentService
from services.cash_shift_service import CashShiftService


def _receipt(db, source_type, source_id, amount, collected_by_id, method="cash", created_at=None):
    """Seed a real source document before calling the production ledger seam."""
    if source_type == "monthly_pass":
        source = db.get(MonthlyPass, int(source_id))
    else:
        source = db.get(ParkingSession, str(source_id))
    if source is None:
        vehicle_type = db.scalars(select(VehicleType)).first()
        if vehicle_type is None:
            vehicle_type = VehicleType(name="Finance fixture vehicle type")
            db.add(vehicle_type)
            db.flush()
        vehicle = db.scalars(select(Vehicle)).first()
        if vehicle is None:
            vehicle = Vehicle(license_plate="51A-100.01", vehicle_type_id=vehicle_type.id)
            db.add(vehicle)
            db.flush()
        if source_type == "monthly_pass":
            customer = db.scalars(select(Customer)).first()
            if customer is None:
                customer = Customer(full_name="Finance fixture", phone_number="0901234567")
                db.add(customer)
                db.flush()
            source = MonthlyPass(id=int(source_id), customer_id=customer.id, vehicle_id=vehicle.id, pass_code=f"LEDGER-{source_id}", price=amount,
                                 start_date=datetime(2026, 1, 1).date(), end_date=datetime(2026, 1, 31).date(), is_active=False)
        else:
            when = created_at or datetime(2026, 1, 1)
            source = ParkingSession(id=str(source_id), vehicle_id=vehicle.id, staff_in_id=collected_by_id, staff_out_id=collected_by_id,
                                    check_in_time=when-timedelta(hours=1), check_out_time=when, status="completed", parking_fee=amount)
        db.add(source)
        db.flush()
    return PaymentService.record_receipt(db, source_type, source_id, amount, collected_by_id, method=method, created_at=created_at)


@pytest.fixture
def manager(db_session):
    role = Role(name="manager")
    db_session.add(role)
    db_session.flush()
    manager = User(role_id=role.id, username="finance_manager", full_name="Quản lý thu ngân", password_hash="unused", is_active=True)
    db_session.add(manager)
    db_session.commit()
    return manager


def headers(user):
    return {"Authorization": "Bearer " + AuthService().create_access_token(user_id=user.id, username=user.username, role=user.role.name)}


def test_receipt_is_idempotent_and_caller_owns_transaction(db_session, test_user):
    first = _receipt(db_session, "monthly_pass", 12, 500_000, test_user.id)
    second = _receipt(db_session, "monthly_pass", 12, 500_000, test_user.id)
    assert first.id == second.id
    assert db_session.query(Payment).count() == 1
    db_session.rollback()
    assert db_session.query(Payment).count() == 0


def test_monthly_payment_is_included_once_and_by_collection_day(db_session, test_user):
    when = datetime(2026, 9, 6, 0, 5)
    _receipt(db_session, "monthly_pass", 12, 500_000, test_user.id, created_at=when)
    _receipt(db_session, "monthly_pass", 12, 500_000, test_user.id, created_at=when)
    result = PaymentService.revenue_breakdown(db_session, when.replace(hour=0, minute=0), when + timedelta(days=1))
    assert result == {"parking_revenue": 0, "monthly_pass_revenue": 500_000, "refunds": 0, "total_revenue": 500_000}


def test_shift_closing_freezes_cash_and_excludes_transfer(db_session, test_user):
    shift = CashShiftService.open_shift(db_session, test_user, opening_cash=100_000)
    _receipt(db_session, "monthly_pass", 12, 500_000, test_user.id)
    _receipt(db_session, "parking_session", "session-transfer", 20_000, test_user.id, method="transfer")
    closed = CashShiftService.close_shift(db_session, shift.id, test_user, counted_cash=590_000)
    assert closed["expected_cash"] == 600_000
    assert closed["difference"] == -10_000
    assert closed["net_revenue"] == 520_000
    assert closed["payment_count"] == 2
    later = _receipt(db_session, "monthly_pass", 13, 30_000, test_user.id)
    assert later.shift_id is None
    with pytest.raises(HTTPException) as err:
        CashShiftService.close_shift(db_session, shift.id, test_user, counted_cash=600_000)
    assert err.value.status_code == 409


def test_refund_is_bounded_idempotent_and_uses_actual_refund_shift(db_session, test_user, manager):
    old_shift = CashShiftService.open_shift(db_session, test_user)
    receipt = _receipt(db_session, "monthly_pass", 12, 500_000, test_user.id)
    CashShiftService.close_shift(db_session, old_shift.id, test_user, counted_cash=500_000)
    new_shift = CashShiftService.open_shift(db_session, manager, opening_cash=500_000)
    payload = {"amount": 200_000, "method": "cash", "reason": "Khách hủy kỳ tiếp theo", "idempotency_key": "refund-once-12"}
    refund = PaymentService.refund(db_session, receipt.id, manager, **payload)
    retried = PaymentService.refund(db_session, receipt.id, manager, **payload)
    assert refund.id == retried.id
    assert refund.shift_id == new_shift.id
    assert CashShiftService.get_summary(db_session, old_shift.id, test_user)["expected_cash"] == 500_000
    assert CashShiftService.get_summary(db_session, new_shift.id, manager)["expected_cash"] == 300_000
    with pytest.raises(HTTPException) as error:
        PaymentService.refund(db_session, receipt.id, manager, **{**payload, "amount": 300_001, "idempotency_key": "refund-too-much"})
    assert error.value.status_code == 409
    assert PaymentService.serialize(db_session, receipt)["refundable_amount"] == 300_000


def test_refund_cannot_reuse_key_with_different_amount(db_session, manager):
    receipt = _receipt(db_session, "monthly_pass", 1, 500_000, manager.id)
    kwargs = {"method": "transfer", "reason": "Điều chỉnh", "idempotency_key": "refund-fixed-key"}
    PaymentService.refund(db_session, receipt.id, manager, amount=100_000, **kwargs)
    with pytest.raises(HTTPException) as error:
        PaymentService.refund(db_session, receipt.id, manager, amount=1, **kwargs)
    assert error.value.status_code == 409


def test_refund_from_previous_day_can_make_today_negative(db_session, manager):
    from core.clock import business_now, day_bounds
    now = business_now()
    receipt = _receipt(db_session, "monthly_pass", 1, 500_000, manager.id, created_at=now - timedelta(days=1))
    PaymentService.refund(db_session, receipt.id, manager, amount=100_000, method="cash", reason="Hủy", idempotency_key="refund-prev-day")
    start, end = day_bounds(now.date())
    report = PaymentService.revenue_breakdown(db_session, start, end)
    assert report["total_revenue"] == -100_000
    assert PaymentService.revenue_by_day(db_session, start, end)[now.date().isoformat()] == -100_000


def test_revenue_bridge_uses_only_unrecorded_sources(db_session, test_user, customer, vehicle):
    day = datetime(2026, 9, 6)
    session = ParkingSession(vehicle_id=vehicle.id, staff_in_id=test_user.id, staff_out_id=test_user.id,
                             check_in_time=day, check_out_time=day + timedelta(hours=1), status="completed", parking_fee=20_000)
    monthly = MonthlyPass(customer_id=customer.id, vehicle_id=vehicle.id, pass_code="BRIDGE-001", price=500_000,
                          start_date=day.date(), end_date=(day+timedelta(days=30)).date(), created_at=day - timedelta(hours=6), is_active=True)
    db_session.add_all([session, monthly])
    db_session.flush()
    assert PaymentService.revenue_breakdown(db_session, day, day+timedelta(days=1))["total_revenue"] == 520_000
    _receipt(db_session, "parking_session", session.id, 20_000, test_user.id, created_at=session.check_out_time)
    _receipt(db_session, "monthly_pass", monthly.id, 500_000, test_user.id, created_at=day + timedelta(hours=1))
    assert PaymentService.revenue_breakdown(db_session, day, day+timedelta(days=1))["total_revenue"] == 520_000


@pytest.mark.parametrize("mutation", ["UPDATE payments SET amount = 0 WHERE id = :id", "DELETE FROM payments WHERE id = :id"])
def test_ledger_is_immutable_at_database(db_session, test_user, mutation):
    payment = _receipt(db_session, "monthly_pass", 12, 500_000, test_user.id)
    db_session.commit()
    with pytest.raises(IntegrityError, match="payment is immutable"):
        db_session.execute(text(mutation), {"id": payment.id})
    db_session.rollback()


def test_database_refund_limit_applies_to_direct_writes(db_session, manager):
    receipt = _receipt(db_session, "monthly_pass", 12, 100, manager.id)
    db_session.commit()
    invalid = Payment(source_type="monthly_pass", source_id="12", kind="refund", amount=101, method="cash",
                      collected_by_id=manager.id, original_payment_id=receipt.id, reason="invalid", idempotency_key="direct-overrefund")
    db_session.add(invalid)
    with pytest.raises(IntegrityError, match="refund exceeds original payment"):
        db_session.flush()
    db_session.rollback()


def test_database_blocks_fake_close_balance_and_closed_shift_payment(db_session, test_user):
    shift = CashShiftService.open_shift(db_session, test_user)
    _receipt(db_session, "monthly_pass", 1, 100, test_user.id)
    db_session.commit()
    with pytest.raises(IntegrityError, match="cash shift close balance invalid"):
        db_session.execute(text("UPDATE cash_shifts SET status='closed', closed_at=:now, counted_cash=0, expected_cash=0, difference=0 WHERE id=:id"), {"now": datetime(2099,1,1), "id": shift.id})
    db_session.rollback()
    CashShiftService.close_shift(db_session, shift.id, test_user, counted_cash=100)
    db_session.commit()
    db_session.add(Payment(source_type="monthly_pass", source_id="1", kind="receipt", amount=100, method="cash",
                           collected_by_id=test_user.id, shift_id=shift.id, idempotency_key="direct-closed"))
    with pytest.raises(IntegrityError, match="payment requires own open shift"):
        db_session.flush()
    db_session.rollback()


def test_staff_sees_own_payments_and_cannot_refund(client, db_session, test_user, manager):
    own = _receipt(db_session, "monthly_pass", 1, 100, test_user.id)
    other = _receipt(db_session, "monthly_pass", 2, 200, manager.id)
    db_session.commit()
    response = client.get("/api/v1/payments", headers=headers(test_user))
    assert response.status_code == 200, response.text
    assert [row["id"] for row in response.json()["items"]] == [own.id]
    assert client.get(f"/api/v1/payments/{other.id}", headers=headers(test_user)).status_code == 403
    assert client.get("/api/v1/payments", headers=headers(manager)).json()["total"] == 2
    response = client.post(f"/api/v1/payments/{own.id}/refund", headers=headers(test_user), json={"amount": 50, "method": "cash", "reason": "Test", "idempotency_key": "api-refund"})
    assert response.status_code == 403


def test_shift_api_open_close_access_and_strict_money(client, db_session, test_user, manager):
    own_headers = headers(test_user)
    assert client.post("/api/v1/cash-shifts", headers=own_headers, json={"opening_cash": 1.5}).status_code == 422
    response = client.post("/api/v1/cash-shifts", headers=own_headers, json={"opening_cash": 100_000})
    assert response.status_code == 201, response.text
    shift_id = response.json()["id"]
    assert client.post("/api/v1/cash-shifts", headers=own_headers, json={"opening_cash": 0}).status_code == 409
    other = CashShiftService.open_shift(db_session, manager)
    db_session.commit()
    assert client.get(f"/api/v1/cash-shifts/{other.id}", headers=own_headers).status_code == 403
    assert client.post(f"/api/v1/cash-shifts/{other.id}/close", headers=own_headers, json={"counted_cash": 0}).status_code == 403
    result = client.post(f"/api/v1/cash-shifts/{shift_id}/close", headers=own_headers, json={"counted_cash": 95_000})
    assert result.status_code == 200, result.text
    assert result.json()["difference"] == -5000


def test_report_keeps_session_average_and_exports_monthly_receipts(client, db_session, manager):
    from core.clock import business_now
    from io import BytesIO
    from openpyxl import load_workbook
    _receipt(db_session, "monthly_pass", 1, 500_000, manager.id)
    db_session.commit()
    response = client.get("/reports/revenue?period=day", headers=headers(manager))
    assert response.status_code == 200, response.text
    report = response.json()
    assert report["total_revenue"] == report["monthly_pass_revenue"] == 500_000
    assert report["average_fee"] == report["total_trips"] == 0
    exported = client.get("/reports/export/xlsx?period=day", headers=headers(manager))
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content))["Tong quan"]
    labels = {row[0]: row[1] for row in sheet.iter_rows(values_only=True)}
    assert labels["Thu từ vé tháng"] == 500_000


def test_database_rejects_orphan_receipt_and_wrong_source_amount(db_session, test_user):
    _receipt(db_session, "monthly_pass", 1, 100, test_user.id)
    db_session.commit()
    for source_id, amount in (("999", 100), ("1", 101)):
        db_session.add(Payment(source_type="monthly_pass", source_id=source_id, kind="receipt", amount=amount,
                               method="cash", collected_by_id=test_user.id, idempotency_key=f"bad-source-{source_id}"))
        with pytest.raises(IntegrityError, match="payment source or amount invalid"):
            db_session.flush()
        db_session.rollback()


def test_database_preserves_paid_parking_source(db_session, test_user):
    payment = _receipt(db_session, "parking_session", "paid-source-proof", 20_000, test_user.id)
    db_session.commit()
    with pytest.raises(IntegrityError, match="paid parking session cannot be deleted"):
        db_session.execute(text("DELETE FROM parking_sessions WHERE id = :id"), {"id": payment.source_id})
    db_session.rollback()


def test_legacy_bridge_does_not_recount_receipt_outside_selected_period(db_session, test_user, customer, vehicle):
    day = datetime(2026, 9, 6)
    monthly = MonthlyPass(customer_id=customer.id, vehicle_id=vehicle.id, pass_code="OTHER-PERIOD", price=500_000,
                          start_date=day.date(), end_date=(day+timedelta(days=30)).date(),
                          created_at=day-timedelta(hours=7), is_active=True)
    db_session.add(monthly)
    db_session.flush()
    assert PaymentService.revenue_breakdown(db_session, day, day+timedelta(days=1))["total_revenue"] == 500_000
    _receipt(db_session, "monthly_pass", monthly.id, 500_000, test_user.id, created_at=day+timedelta(days=1))
    assert PaymentService.revenue_breakdown(db_session, day, day+timedelta(days=1))["total_revenue"] == 0
    assert PaymentService.revenue_breakdown(db_session, day+timedelta(days=1), day+timedelta(days=2))["total_revenue"] == 500_000
